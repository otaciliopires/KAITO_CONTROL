from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime

import openpyxl
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.messages import constants
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Max, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from ativos.models import Obras
from .models import EntradaDiesel, EquipamentoObra, SaidaDiesel, TanqueObra


def _get_obra(request, obra_id):
    obra = get_object_or_404(Obras, id=obra_id, status='E')
    if not (request.user.is_superuser or obra.usuario.filter(id=request.user.id).exists()):
        return None, HttpResponse('<h1>Acesso negado</h1>')
    return obra, None


def _soma(queryset, campo):
    return queryset.aggregate(total=Sum(campo))['total'] or 0


def _recalcular_estoque(tanque):
    total_entradas = _soma(EntradaDiesel.objects.filter(obra=tanque.obra), 'quantidade')
    total_saidas = _soma(SaidaDiesel.objects.filter(obra=tanque.obra), 'litros')
    tanque.estoque = total_entradas - total_saidas
    tanque.save()


def _parse_data(valor, padrao):
    if not valor:
        return padrao
    return datetime.strptime(valor, '%Y-%m-%d').date()


MESES_PT = {
    1: 'Jan', 2: 'Fev', 3: 'Mar', 4: 'Abr', 5: 'Mai', 6: 'Jun',
    7: 'Jul', 8: 'Ago', 9: 'Set', 10: 'Out', 11: 'Nov', 12: 'Dez',
}


def _consumo_mensal_e_diario(obra):
    consumo_mensal = defaultdict(float)
    consumo_diario = defaultdict(lambda: defaultdict(float))

    for saida in SaidaDiesel.objects.filter(obra=obra).only('data', 'litros'):
        chave_mes = saida.data.strftime('%Y-%m')
        consumo_mensal[chave_mes] += saida.litros
        consumo_diario[chave_mes][saida.data.day] += saida.litros

    meses_ordenados = sorted(consumo_mensal.keys())

    dados_diarios = {}
    for chave_mes in meses_ordenados:
        ano, mes_num = int(chave_mes[:4]), int(chave_mes[5:7])
        dias_no_mes = monthrange(ano, mes_num)[1]
        dados_diarios[chave_mes] = {
            'label': f"{MESES_PT[mes_num]}/{ano}",
            'dias': list(range(1, dias_no_mes + 1)),
            'valores': [round(consumo_diario[chave_mes].get(dia, 0), 2) for dia in range(1, dias_no_mes + 1)],
        }

    return {
        'meses_chaves': meses_ordenados,
        'meses_labels': [dados_diarios[m]['label'] for m in meses_ordenados],
        'meses_valores': [round(consumo_mensal[m], 2) for m in meses_ordenados],
        'dados_diarios': dados_diarios,
    }


@login_required(login_url='/auth/login/')
def home(request, obra_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    tanque, _created = TanqueObra.objects.get_or_create(obra=obra)
    _recalcular_estoque(tanque)

    consumo_total = _soma(SaidaDiesel.objects.filter(obra=obra), 'litros')
    entradas_total = _soma(EntradaDiesel.objects.filter(obra=obra), 'quantidade')

    ranking = []
    for equipamento in EquipamentoObra.objects.filter(obra=obra):
        saidas_equip = SaidaDiesel.objects.filter(equipamento=equipamento)
        litros = _soma(saidas_equip, 'litros')
        if not litros:
            continue
        percurso = sum(s.contador_fim - s.contador_inicio for s in saidas_equip)
        media = (litros / percurso) if percurso else 0
        ranking.append({
            'equipamento': equipamento,
            'litros': litros,
            'percurso': percurso,
            'media': media,
        })
    ranking.sort(key=lambda item: item['litros'], reverse=True)

    consumo_grafico = _consumo_mensal_e_diario(obra)

    return render(request, 'cdo_home.html', {
        'obra': obra,
        'tanque': tanque,
        'consumo_total': consumo_total,
        'entradas_total': entradas_total,
        'ranking': ranking,
        'consumo_meses_labels': consumo_grafico['meses_labels'],
        'consumo_meses_valores': consumo_grafico['meses_valores'],
        'consumo_meses_chaves': consumo_grafico['meses_chaves'],
        'consumo_dados_diarios': consumo_grafico['dados_diarios'],
    })


@login_required(login_url='/auth/login/')
def operacoes(request, obra_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    tanque, _created = TanqueObra.objects.get_or_create(obra=obra)

    if request.method == 'POST':
        form_entradas = request.POST.get('form_entradas')
        form_saidas = request.POST.get('form_saidas')

        if form_entradas:
            fornecedor = request.POST.get('fornecedor')
            nota_fiscal = request.POST.get('nota_fiscal')
            data_nf = request.POST.get('data_nf')
            data_entrega = request.POST.get('data_entrega')
            valor_litro = request.POST.get('valor_litro')
            quantidade = request.POST.get('quantidade')
            descricao = request.POST.get('descricao', '')

            valor_total = float(valor_litro) * float(quantidade)
            num_entrada = EntradaDiesel.objects.filter(obra=obra).aggregate(Max('numero'))['numero__max']
            num_entrada = (num_entrada or 0) + 1

            EntradaDiesel.objects.create(
                obra=obra,
                numero=num_entrada,
                nota_fiscal=nota_fiscal,
                fornecedor=fornecedor,
                data_nf=data_nf,
                data_entrega=data_entrega,
                quantidade=quantidade,
                preco_unitario=valor_litro,
                preco_total=valor_total,
                colaborador=request.user,
                descricao=descricao,
            )
            _recalcular_estoque(tanque)
            messages.add_message(request, constants.SUCCESS, "Entrada lançada com sucesso!")
            return redirect('cdo_operacoes', obra_id=obra.id)

        if form_saidas:
            equipamento_id = request.POST.get('equipamento')
            equipamento = get_object_or_404(EquipamentoObra, id=equipamento_id, obra=obra)

            contador_inicial = float(request.POST.get('contador_inicial'))
            contador_final = float(request.POST.get('contador_final'))
            litros = request.POST.get('litros')
            operador = request.POST.get('operador')
            data = request.POST.get('data')
            observacao = request.POST.get('observacao', '')

            num_saida = SaidaDiesel.objects.filter(obra=obra).aggregate(Max('numero'))['numero__max']
            num_saida = (num_saida or 0) + 1

            SaidaDiesel.objects.create(
                obra=obra,
                numero=num_saida,
                equipamento=equipamento,
                litros=litros,
                contador_inicio=contador_inicial,
                contador_fim=contador_final,
                data=data,
                operador=operador,
                colaborador=request.user,
                observacao=observacao,
            )
            equipamento.contador_atual = contador_final
            equipamento.save()
            _recalcular_estoque(tanque)
            messages.add_message(request, constants.SUCCESS, "Saída lançada com sucesso!")
            return redirect('cdo_operacoes', obra_id=obra.id)

    equipamentos = EquipamentoObra.objects.filter(obra=obra, ativo=True).order_by('prefixo')
    return render(request, 'cdo_operacoes.html', {
        'obra': obra,
        'tanque': tanque,
        'equipamentos': equipamentos,
    })


@login_required(login_url='/auth/login/')
def entradas(request, obra_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    data_inicio = _parse_data(request.GET.get('data_inicio'), date(2020, 1, 1))
    data_fim = _parse_data(request.GET.get('data_fim'), date.today())

    entradas_qs = EntradaDiesel.objects.filter(
        obra=obra, data_entrega__range=[data_inicio, data_fim]
    ).order_by('-numero')
    total_entradas = _soma(entradas_qs, 'quantidade')

    paginator = Paginator(entradas_qs, 50)
    entradas_page = paginator.get_page(request.GET.get('page'))

    querystring = request.GET.copy()
    querystring.pop('page', None)
    querystring = querystring.urlencode()

    return render(request, 'cdo_entradas.html', {
        'obra': obra,
        'entradas': entradas_page,
        'total_entradas': total_entradas,
        'data_inicio_raw': request.GET.get('data_inicio', ''),
        'data_fim_raw': request.GET.get('data_fim', ''),
        'querystring': querystring,
    })


@login_required(login_url='/auth/login/')
def entrada_editar(request, obra_id, entrada_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    entrada = get_object_or_404(EntradaDiesel, id=entrada_id, obra=obra)
    tanque, _created = TanqueObra.objects.get_or_create(obra=obra)

    if request.method == 'POST':
        entrada.nota_fiscal = request.POST.get('nota_fiscal')
        entrada.fornecedor = request.POST.get('fornecedor')
        entrada.data_nf = request.POST.get('data_nf')
        entrada.data_entrega = request.POST.get('data_entrega')
        entrada.quantidade = float(request.POST.get('quantidade'))
        entrada.preco_unitario = float(request.POST.get('valor_litro'))
        entrada.preco_total = entrada.quantidade * entrada.preco_unitario
        entrada.descricao = request.POST.get('descricao', '')
        entrada.save()
        _recalcular_estoque(tanque)
        messages.add_message(request, constants.SUCCESS, "Entrada atualizada com sucesso!")
        return redirect('cdo_entradas', obra_id=obra.id)

    return render(request, 'cdo_entrada_editar.html', {'obra': obra, 'entrada': entrada})


@login_required(login_url='/auth/login/')
def entrada_excluir(request, obra_id, entrada_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    entrada = get_object_or_404(EntradaDiesel, id=entrada_id, obra=obra)
    tanque, _created = TanqueObra.objects.get_or_create(obra=obra)

    if request.method == 'POST':
        entrada.delete()
        _recalcular_estoque(tanque)
        messages.add_message(request, constants.SUCCESS, "Entrada excluída com sucesso!")

    return redirect('cdo_entradas', obra_id=obra.id)


@login_required(login_url='/auth/login/')
def saidas(request, obra_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    equipamentos = EquipamentoObra.objects.filter(obra=obra).order_by('prefixo')

    data_inicio = _parse_data(request.GET.get('data_inicio'), date(2020, 1, 1))
    data_fim = _parse_data(request.GET.get('data_fim'), date.today())
    filtro_equipamento = request.GET.getlist('equipamento')

    saidas_qs = SaidaDiesel.objects.filter(obra=obra, data__range=[data_inicio, data_fim]).order_by('-numero')
    if filtro_equipamento:
        saidas_qs = saidas_qs.filter(equipamento_id__in=filtro_equipamento)

    total_saidas = _soma(saidas_qs, 'litros')

    paginator = Paginator(saidas_qs, 50)
    saidas_page = paginator.get_page(request.GET.get('page'))

    querystring = request.GET.copy()
    querystring.pop('page', None)
    querystring = querystring.urlencode()

    return render(request, 'cdo_saidas.html', {
        'obra': obra,
        'saidas': saidas_page,
        'equipamentos': equipamentos,
        'total_saidas': total_saidas,
        'equipamento_ids_selecionados': filtro_equipamento,
        'data_inicio_raw': request.GET.get('data_inicio', ''),
        'data_fim_raw': request.GET.get('data_fim', ''),
        'querystring': querystring,
    })


@login_required(login_url='/auth/login/')
def saida_editar(request, obra_id, saida_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    saida = get_object_or_404(SaidaDiesel, id=saida_id, obra=obra)
    tanque, _created = TanqueObra.objects.get_or_create(obra=obra)
    equipamentos = EquipamentoObra.objects.filter(obra=obra).order_by('prefixo')

    if request.method == 'POST':
        equipamento_id = request.POST.get('equipamento')
        saida.equipamento = get_object_or_404(EquipamentoObra, id=equipamento_id, obra=obra)
        saida.litros = float(request.POST.get('litros'))
        saida.contador_inicio = float(request.POST.get('contador_inicial'))
        saida.contador_fim = float(request.POST.get('contador_final'))
        saida.data = request.POST.get('data')
        saida.operador = request.POST.get('operador')
        saida.observacao = request.POST.get('observacao', '')
        saida.save()
        _recalcular_estoque(tanque)
        messages.add_message(request, constants.SUCCESS, "Saída atualizada com sucesso!")
        return redirect('cdo_saidas', obra_id=obra.id)

    return render(request, 'cdo_saida_editar.html', {'obra': obra, 'saida': saida, 'equipamentos': equipamentos})


@login_required(login_url='/auth/login/')
def saida_excluir(request, obra_id, saida_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    saida = get_object_or_404(SaidaDiesel, id=saida_id, obra=obra)
    tanque, _created = TanqueObra.objects.get_or_create(obra=obra)

    if request.method == 'POST':
        saida.delete()
        _recalcular_estoque(tanque)
        messages.add_message(request, constants.SUCCESS, "Saída excluída com sucesso!")

    return redirect('cdo_saidas', obra_id=obra.id)


@login_required(login_url='/auth/login/')
def equipamentos(request, obra_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    if request.method == 'POST':
        prefixo = request.POST.get('prefixo')
        descricao = request.POST.get('descricao', '')
        tipo = request.POST.get('tipo')
        unidade_medicao = request.POST.get('unidade_medicao')
        contador_atual = request.POST.get('contador_atual') or 0

        if EquipamentoObra.objects.filter(obra=obra, prefixo=prefixo).exists():
            messages.add_message(
                request, constants.ERROR,
                f"Já existe um equipamento com o prefixo '{prefixo}' nesta obra.",
            )
        else:
            EquipamentoObra.objects.create(
                obra=obra,
                prefixo=prefixo,
                descricao=descricao,
                tipo=tipo,
                unidade_medicao=unidade_medicao,
                contador_atual=contador_atual,
            )
            messages.add_message(request, constants.SUCCESS, "Equipamento cadastrado com sucesso!")
        return redirect('cdo_equipamentos', obra_id=obra.id)

    equipamentos_qs = EquipamentoObra.objects.filter(obra=obra).order_by('prefixo')
    return render(request, 'cdo_equipamentos.html', {'obra': obra, 'equipamentos': equipamentos_qs})


@login_required(login_url='/auth/login/')
def equipamento_editar(request, obra_id, equip_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    equipamento = get_object_or_404(EquipamentoObra, id=equip_id, obra=obra)

    if request.method == 'POST':
        equipamento.prefixo = request.POST.get('prefixo')
        equipamento.descricao = request.POST.get('descricao', '')
        equipamento.tipo = request.POST.get('tipo')
        equipamento.unidade_medicao = request.POST.get('unidade_medicao')
        equipamento.contador_atual = request.POST.get('contador_atual') or 0
        equipamento.ativo = bool(request.POST.get('ativo'))
        equipamento.save()
        messages.add_message(request, constants.SUCCESS, "Equipamento atualizado com sucesso!")
        return redirect('cdo_equipamentos', obra_id=obra.id)

    return render(request, 'cdo_equipamento_editar.html', {'obra': obra, 'equipamento': equipamento})


def _numero_ou_texto(valor):
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return valor
    return f"'{valor}'" if valor is not None else "(vazio)"


def _validar_entradas_excel(sheet):
    validas = []
    erros = []
    for linha_num, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
        ent = row[:8]
        if ent[0] is None:
            break

        linha_erros = []
        if not isinstance(ent[1], (int, float)):
            linha_erros.append(f"quantidade inválida ({_numero_ou_texto(ent[1])})")
        if not ent[2]:
            linha_erros.append("fornecedor não informado")
        if not ent[3]:
            linha_erros.append("nota fiscal não informada")
        if ent[4] is None:
            linha_erros.append("data da nota fiscal não informada")
        if not isinstance(ent[5], (int, float)):
            linha_erros.append(f"preço unitário inválido ({_numero_ou_texto(ent[5])})")
        if not isinstance(ent[6], (int, float)):
            linha_erros.append(f"preço total inválido ({_numero_ou_texto(ent[6])})")

        if linha_erros:
            erros.append(f"ENTRADAS, linha {linha_num}: " + "; ".join(linha_erros))
        else:
            validas.append({'dados': ent})

    return validas, erros


def _validar_saidas_excel(sheet, obra):
    validas = []
    erros = []
    for linha_num, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
        sds = row[:7]
        if sds[0] is None:
            break

        linha_erros = []
        equipamento = EquipamentoObra.objects.filter(obra=obra, prefixo=sds[1]).first()
        if equipamento is None:
            linha_erros.append(f"equipamento {_numero_ou_texto(sds[1])} não encontrado nesta obra")
        if not isinstance(sds[2], (int, float)):
            linha_erros.append(f"contador inicial inválido ({_numero_ou_texto(sds[2])})")
        if not isinstance(sds[3], (int, float)):
            linha_erros.append(f"contador final inválido ({_numero_ou_texto(sds[3])})")
        if isinstance(sds[2], (int, float)) and isinstance(sds[3], (int, float)) and sds[3] < sds[2]:
            linha_erros.append("contador final menor que o contador inicial")
        if not isinstance(sds[4], (int, float)):
            linha_erros.append(f"litros inválido ({_numero_ou_texto(sds[4])})")
        if not sds[5]:
            linha_erros.append("operador não informado")

        if linha_erros:
            erros.append(f"SAIDAS, linha {linha_num}: " + "; ".join(linha_erros))
        else:
            validas.append({'equipamento': equipamento, 'dados': sds})

    return validas, erros


@login_required(login_url='/auth/login/')
def importar_excel(request, obra_id):
    obra, negado = _get_obra(request, obra_id)
    if negado:
        return negado

    tanque, _created = TanqueObra.objects.get_or_create(obra=obra)

    if request.method != 'POST':
        return redirect('cdo_operacoes', obra_id=obra.id)

    if 'excel' not in request.FILES:
        messages.add_message(request, constants.ERROR, "Nenhum arquivo enviado.")
        return redirect('cdo_operacoes', obra_id=obra.id)

    try:
        workbook = openpyxl.load_workbook(request.FILES['excel'])
    except Exception:
        messages.add_message(
            request, constants.ERROR,
            "Não foi possível ler o arquivo. Verifique se é um .xlsx válido.",
        )
        return redirect('cdo_operacoes', obra_id=obra.id)

    abas_esperadas = ['ENTRADAS', 'SAIDAS']
    abas_faltando = [a for a in abas_esperadas if a not in workbook.sheetnames]
    if abas_faltando:
        messages.add_message(
            request, constants.ERROR,
            f"Planilha inválida: aba(s) não encontrada(s): {', '.join(abas_faltando)}",
        )
        return redirect('cdo_operacoes', obra_id=obra.id)

    entradas_validas, erros_entradas = _validar_entradas_excel(workbook['ENTRADAS'])
    saidas_validas, erros_saidas = _validar_saidas_excel(workbook['SAIDAS'], obra)

    erros = erros_entradas + erros_saidas
    if erros:
        messages.add_message(
            request, constants.ERROR,
            f"Importação cancelada: {len(erros)} erro(s) encontrado(s). "
            "Nenhum lançamento foi feito. Corrija a planilha e envie novamente.",
        )
        for erro in erros:
            messages.add_message(request, constants.ERROR, erro)
        return redirect('cdo_operacoes', obra_id=obra.id)

    if not entradas_validas and not saidas_validas:
        messages.add_message(request, constants.WARNING, "Nenhum lançamento encontrado na planilha.")
        return redirect('cdo_operacoes', obra_id=obra.id)

    with transaction.atomic():
        for item in entradas_validas:
            ent = item['dados']
            num_entrada = EntradaDiesel.objects.filter(obra=obra).aggregate(Max('numero'))['numero__max']
            num_entrada = (num_entrada or 0) + 1
            EntradaDiesel.objects.create(
                obra=obra,
                numero=num_entrada,
                data_entrega=ent[0],
                quantidade=ent[1],
                fornecedor=ent[2],
                nota_fiscal=ent[3],
                data_nf=ent[4],
                preco_unitario=ent[5],
                preco_total=ent[6],
                colaborador=request.user,
                descricao=(ent[7] or "") if len(ent) > 7 else "",
            )

        for item in saidas_validas:
            sds = item['dados']
            equipamento = item['equipamento']
            num_saida = SaidaDiesel.objects.filter(obra=obra).aggregate(Max('numero'))['numero__max']
            num_saida = (num_saida or 0) + 1
            SaidaDiesel.objects.create(
                obra=obra,
                numero=num_saida,
                equipamento=equipamento,
                data=sds[0],
                contador_inicio=sds[2],
                contador_fim=sds[3],
                litros=sds[4],
                operador=sds[5],
                colaborador=request.user,
                observacao=(sds[6] or "") if len(sds) > 6 else "",
            )
            equipamento.contador_atual = sds[3]
            equipamento.save()

    _recalcular_estoque(tanque)

    messages.add_message(
        request, constants.SUCCESS,
        f"Arquivo importado com sucesso! {len(entradas_validas)} entrada(s), "
        f"{len(saidas_validas)} saída(s) lançada(s).",
    )
    return redirect('cdo_operacoes', obra_id=obra.id)
